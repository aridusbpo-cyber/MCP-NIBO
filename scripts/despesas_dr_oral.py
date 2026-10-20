"""Monta o bloco de despesas do Nibo abaixo do RESUMO das planilhas
"DR Oral Clinic <unidade> - Controle Anual 2026".

Regras (aprovadas pela cliente em 20/10/2026):
- unidade = centro de custo do Nibo (CAJAZEIRAS-PB / TRIUNFO-PB); demais CCs ficam fora
- área = categoria do Nibo (MAPA_AREAS); Prótese + Implante numa linha só
- "Materiais Odontológicos" (2.405) em linha própria
- "Despesas gerais" para o resto das despesas operacionais
- data = data do pagamento; transferências e movimentações não operacionais excluídas
- quadro Receita x Despesa x Resultado; receita = linha "TOTAL DO MÊS" do RESUMO

Uso:
    python despesas_dr_oral.py --ano 2026 --cajazeiras caj.xlsx --triunfo tri.xlsx [--saida-dir DIR]

Token: o proxy do ambiente injeta o ApiToken; fora dele, defina NIBO_API_TOKEN.
"""
import argparse
import datetime as dt
import json
import os
import re
import urllib.parse
import urllib.request
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

BASE = "https://api.nibo.com.br/empresas/v1"
FUSO_BR = dt.timezone(dt.timedelta(hours=-3))

UNIDADES = {"cajazeiras": "CAJAZEIRAS-PB", "triunfo": "TRIUNFO-PB"}

# Ordem = ordem das linhas no bloco. Casamento pelo código da categoria (prefixo do nome).
MAPA_AREAS = [
    ("HOF", ["2.402", "4.108"]),
    ("CLÍNICO", ["2.409"]),
    ("ORTO", ["2.411"]),
    ("PRÓTESE + IMPLANTE", ["2.407", "2.410", "Compra de Material - Implantes"]),
    ("LENTES/REC. ESTÉTICA", ["2.408"]),
    ("MATERIAIS ODONTOLÓGICOS", ["2.405"]),
]
LINHA_GERAIS = "DESPESAS GERAIS"

# Fora do quadro: não operacional, devoluções de venda e ajustes.
EXCLUIR = ["5.", "4.201", "2.201", "2.202", "Capitalização Sicredi", "AJUSTE SALDO CONTA"]

MARCADOR = "DESPESAS NIBO POR ÁREA"
MESES = ["JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT", "NOV", "DEZ"]


def _casa(nome: str, chaves) -> bool:
    return any(nome.strip().startswith(k) for k in chaves)


def area_da_categoria(nome: str):
    if _casa(nome, EXCLUIR):
        return None
    for area, chaves in MAPA_AREAS:
        if _casa(nome, chaves):
            return area
    return LINHA_GERAIS


def buscar_pagamentos(ano: int) -> list:
    headers = {"Accept": "application/json"}
    if os.environ.get("NIBO_API_TOKEN"):
        headers["ApiToken"] = os.environ["NIBO_API_TOKEN"]
    filtro = f"date ge {ano}-01-01T00:00:00Z and date lt {ano + 1}-01-01T00:00:00Z"
    itens, skip = [], 0
    while True:
        q = urllib.parse.urlencode({"$filter": filtro, "$top": 500, "$skip": skip})
        req = urllib.request.Request(f"{BASE}/payments?{q}", headers=headers)
        with urllib.request.urlopen(req, timeout=60) as r:
            pagina = json.load(r)["items"]
        itens += pagina
        if len(pagina) < 500:
            return itens
        skip += 500


def agregar(pagamentos: list):
    """{centro_de_custo: {area: [12 meses]}} + lista de pendências (sem CC de unidade)."""
    tot = defaultdict(lambda: defaultdict(lambda: [0.0] * 12))
    pendencias = []
    for p in pagamentos:
        if p.get("isTransfer"):
            continue
        mes = int(p["date"][5:7]) - 1
        ccs = p.get("costCenters") or []
        for c in p["categories"]:
            area = area_da_categoria(c["categoryName"])
            if area is None:
                continue
            # categoria de entrada dentro de um pagamento (ex.: Descontos Recebidos) reduz a despesa
            valor = -c["value"] if c.get("type") == "in" else c["value"]
            if not ccs:
                pendencias.append((p["date"][:10], c["categoryName"], "(sem centro de custo)",
                                   valor, p.get("description") or ""))
            for cc in ccs:
                nome_cc = cc["costCenterDescription"]
                parte = valor * cc["percent"] / 100
                if nome_cc in UNIDADES.values():
                    tot[nome_cc][area][mes] += parte
                elif nome_cc == "Não Identificado":
                    pendencias.append((p["date"][:10], c["categoryName"], nome_cc, parte,
                                       p.get("description") or ""))
                # SÓCIOS e OBRA: fora do quadro, por decisão da cliente
    return tot, pendencias


def _linha_total_resumo(ws) -> int:
    for r in range(1, 60):
        if str(ws.cell(r, 1).value or "").strip().upper() == "TOTAL DO MÊS":
            return r
    raise ValueError("Linha 'TOTAL DO MÊS' não encontrada no RESUMO")


def escrever_bloco(caminho_in: Path, caminho_out: Path, valores: dict, atualizado_em: str):
    wb = openpyxl.load_workbook(caminho_in)
    ws = wb["RESUMO"]
    lin_receita = _linha_total_resumo(ws)

    # Limpa bloco anterior (do marcador para baixo) para a atualização ser idempotente.
    inicio = lin_receita + 3
    for r in range(lin_receita + 1, ws.max_row + 1):
        if str(ws.cell(r, 1).value or "").startswith(MARCADOR):
            inicio = r
            break
    for rng in [m for m in ws.merged_cells.ranges if m.min_row >= inicio]:
        ws.unmerge_cells(str(rng))
    for r in range(inicio, ws.max_row + 1):
        for c in range(1, 15):
            cel = ws.cell(r, c)
            cel.value = None
            cel.style = "Normal"

    azul = PatternFill("solid", fgColor="1F4E78")
    cinza = PatternFill("solid", fgColor="D9E1F2")
    branco_neg = Font(bold=True, color="FFFFFF")
    neg = Font(bold=True)
    fino = Side(style="thin", color="BFBFBF")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    moeda = '"R$" #,##0.00;[Red]-"R$" #,##0.00'

    def cabecalho(r, rotulo):
        ws.cell(r, 1, rotulo)
        for i, m in enumerate(MESES):
            ws.cell(r, 2 + i, m)
        ws.cell(r, 14, "TOTAL ANO")
        for c in range(1, 15):
            ws.cell(r, c).fill, ws.cell(r, c).font = cinza, neg
            ws.cell(r, c).border = borda
            ws.cell(r, c).alignment = Alignment(horizontal="center")

    def formatar_linha(r, negrito=False):
        for c in range(1, 15):
            cel = ws.cell(r, c)
            cel.border = borda
            if c > 1:
                cel.number_format = moeda
            if negrito:
                cel.font = neg

    r = inicio
    ws.cell(r, 1, f"{MARCADOR} (Nibo, data de pagamento) — atualizado em {atualizado_em}")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=14)
    ws.cell(r, 1).fill, ws.cell(r, 1).font = azul, branco_neg
    r += 1
    cabecalho(r, "ÁREA")
    r += 1
    prim = r
    for area in [a for a, _ in MAPA_AREAS] + [LINHA_GERAIS]:
        ws.cell(r, 1, area)
        for i, v in enumerate(valores.get(area, [0.0] * 12)):
            ws.cell(r, 2 + i, round(v, 2))
        ws.cell(r, 14, f"=SUM(B{r}:M{r})")
        formatar_linha(r)
        r += 1
    lin_desp = r
    ws.cell(r, 1, "TOTAL DESPESAS")
    for c in range(2, 15):
        col = openpyxl.utils.get_column_letter(c)
        ws.cell(r, c, f"=SUM({col}{prim}:{col}{r - 1})")
    formatar_linha(r, True)

    r += 3
    ws.cell(r, 1, "RECEITA x DESPESA x RESULTADO")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=14)
    ws.cell(r, 1).fill, ws.cell(r, 1).font = azul, branco_neg
    r += 1
    cabecalho(r, "")
    r += 1
    lr, ld, lres = r, r + 1, r + 2
    ws.cell(lr, 1, "RECEITA")
    ws.cell(ld, 1, "DESPESA")
    ws.cell(lres, 1, "RESULTADO")
    ws.cell(lres + 1, 1, "MARGEM %")
    for c in range(2, 15):
        col = openpyxl.utils.get_column_letter(c)
        ws.cell(lr, c, f"={col}{lin_receita}")
        ws.cell(ld, c, f"={col}{lin_desp}")
        ws.cell(lres, c, f"={col}{lr}-{col}{ld}")
        ws.cell(lres + 1, c, f'=IFERROR({col}{lres}/{col}{lr},"")')
    for x in (lr, ld):
        formatar_linha(x)
    formatar_linha(lres, True)
    formatar_linha(lres + 1)
    for c in range(2, 15):
        ws.cell(lres + 1, c).number_format = "0.0%"

    wb.save(caminho_out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ano", type=int, default=dt.datetime.now(FUSO_BR).year)
    ap.add_argument("--cajazeiras", type=Path, required=True)
    ap.add_argument("--triunfo", type=Path, required=True)
    ap.add_argument("--saida-dir", type=Path, default=Path("."))
    a = ap.parse_args()

    tot, pend = agregar(buscar_pagamentos(a.ano))
    agora = dt.datetime.now(FUSO_BR).strftime("%d/%m/%Y %H:%M")
    a.saida_dir.mkdir(parents=True, exist_ok=True)
    resumo = {}
    for chave, arq in (("cajazeiras", a.cajazeiras), ("triunfo", a.triunfo)):
        cc = UNIDADES[chave]
        escrever_bloco(arq, a.saida_dir / arq.name, tot[cc], agora)
        resumo[cc] = {area: round(sum(m), 2) for area, m in tot[cc].items()}
    resumo["pendencias"] = [
        {"data": d, "categoria": c, "centro_de_custo": cc, "valor": round(v, 2), "descricao": desc}
        for d, c, cc, v, desc in sorted(pend)
    ]
    print(json.dumps(resumo, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
